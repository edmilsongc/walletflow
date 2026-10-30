from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape

from flask import Blueprint, jsonify, request, send_file, session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from src.database.connection import get_connection

exports_bp = Blueprint("exports", __name__, url_prefix="/api")

EXCEL_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TABLE_HEADERS = {
    "receivables": [
        "Descrição", "Categoria", "Valor (R$)", "Vencimento", "Status", "Frequência", "Observações"
    ],
    "payables": [
        "Descrição", "Categoria", "Fornecedor", "Valor (R$)", "Vencimento", "Status",
        "Forma de pagamento", "Observações",
    ],
}


def _parse_date(value):
    if value is None:
        return None
    try:
        parsed = date.fromisoformat(value)
    except ValueError:
        raise ValueError("invalid_date") from None
    if parsed.isoformat() != value:
        raise ValueError("invalid_date")
    return parsed


def _load_records(user_id, start_date, end_date):
    filters = ["user_id = %s"]
    parameters = [user_id]
    if start_date:
        filters.append("due_date >= %s")
        parameters.append(start_date)
    if end_date:
        filters.append("due_date <= %s")
        parameters.append(end_date)
    where_clause = " AND ".join(filters)

    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            f"""
            SELECT description, category, amount, due_date, status, frequency, notes
            FROM accounts_receivable
            WHERE {where_clause}
            ORDER BY due_date, description
            """,
            parameters,
        )
        receivables = cursor.fetchall()
        cursor.execute(
            f"""
            SELECT description, category, payable_supplier, amount, due_date, status, payable_method, notes
            FROM accounts_payable
            WHERE {where_clause}
            ORDER BY due_date, description
            """,
            parameters,
        )
        payables = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()

    return receivables, payables


def _amount(value):
    return Decimal(str(value or 0))


def _summary_rows(receivables, payables):
    rows = []
    sections = (
        ("A receber", receivables, 2, 4),
        ("A pagar", payables, 3, 5),
    )
    for label, records, amount_index, status_index in sections:
        grouped = defaultdict(lambda: [0, Decimal("0")])
        for record in records:
            status = record[status_index]
            grouped[status or "Sem status"][0] += 1
            grouped[status or "Sem status"][1] += _amount(record[amount_index])
        for status, (count, total) in sorted(grouped.items()):
            rows.append([label, status, count, total])
        total = sum(
            (_amount(record[amount_index]) for record in records),
            Decimal("0"),
        )
        rows.append([label, "Total", len(records), total])
    return rows


def _excel_safe_text(value):
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return f"'{text}"
    return text


def _add_excel_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6F5C")
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append([
            _excel_safe_text(value) if isinstance(value, str) or value is None else value
            for value in row
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, column in enumerate(sheet.columns, start=1):
        longest = max((len(str(cell.value or "")) for cell in column), default=0)
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(longest + 2, 14), 42)
    return sheet


def _create_xlsx(receivables, payables, include_transactions):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    summary.append(["Tipo", "Status", "Quantidade", "Total (R$)"])
    for row in _summary_rows(receivables, payables):
        summary.append(row)
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6F5C")
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    for column_index, column in enumerate(summary.columns, start=1):
        longest = max((len(str(cell.value or "")) for cell in column), default=0)
        summary.column_dimensions[get_column_letter(column_index)].width = min(max(longest + 2, 16), 36)
    for cell in summary["D"][1:]:
        cell.number_format = '"R$" #,##0.00'

    if include_transactions:
        _add_excel_sheet(workbook, "A receber", TABLE_HEADERS["receivables"], receivables)
        _add_excel_sheet(workbook, "A pagar", TABLE_HEADERS["payables"], payables)
        for title, amount_column in (("A receber", "C"), ("A pagar", "D")):
            for cell in workbook[title][amount_column][1:]:
                cell.number_format = '"R$" #,##0.00'

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _brl(value):
    return f"R$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _pdf_table(headers, rows, widths, style):
    content = [[Paragraph(escape(str(value)), style) for value in headers]]
    content.extend([
        [Paragraph(escape(str(value if value is not None else "")), style) for value in row]
        for row in rows
    ])
    table = Table(content, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F6F5C")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9E2E7")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F7F6")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    return table


def _create_pdf(receivables, payables, include_transactions, start_date, end_date):
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title="Relatório WalletFlow",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="WalletFlowBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="WalletFlowHeading",
        parent=styles["Heading2"],
        textColor=colors.HexColor("#1F6F5C"),
        spaceBefore=10,
        spaceAfter=6,
    ))
    elements = [
        Paragraph("Relatório financeiro WalletFlow", styles["Title"]),
        Paragraph(
            f"Período por vencimento: {start_date.isoformat() if start_date else 'sem limite'} a {end_date.isoformat() if end_date else 'sem limite'}",
            styles["BodyText"],
        ),
        Spacer(1, 8 * mm),
        Paragraph("Resumo", styles["WalletFlowHeading"]),
    ]
    summary = [
        [row[0], row[1], str(row[2]), _brl(row[3])]
        for row in _summary_rows(receivables, payables)
    ]
    elements.append(
        _pdf_table(
            ["Tipo", "Status", "Quantidade", "Total"],
            summary,
            [45 * mm, 65 * mm, 35 * mm, 45 * mm],
            styles["WalletFlowBody"],
        )
    )
    if include_transactions:
        elements.extend([PageBreak(), Paragraph("Contas a receber", styles["WalletFlowHeading"])])
        receivable_rows = [
            [record[0], record[1], _brl(_amount(record[2])), record[3], record[4], record[5], record[6]]
            for record in receivables
        ]
        elements.append(_pdf_table(
            TABLE_HEADERS["receivables"],
            receivable_rows,
            [48 * mm, 30 * mm, 27 * mm, 25 * mm, 26 * mm, 27 * mm, 70 * mm],
            styles["WalletFlowBody"],
        ))
        elements.extend([PageBreak(), Paragraph("Contas a pagar", styles["WalletFlowHeading"])])
        payable_rows = [
            [record[0], record[1], record[2], _brl(_amount(record[3])), record[4], record[5], record[6], record[7]]
            for record in payables
        ]
        elements.append(_pdf_table(
            TABLE_HEADERS["payables"],
            payable_rows,
            [42 * mm, 25 * mm, 35 * mm, 25 * mm, 23 * mm, 25 * mm, 35 * mm, 50 * mm],
            styles["WalletFlowBody"],
        ))
    document.build(elements)
    output.seek(0)
    return output


@exports_bp.route("/exports/dashboard", methods=["GET"])
def export_dashboard():
    user_id = session.get("user_id")
    if user_id is None:
        return jsonify({"message": "authentication_required"}), 401

    file_format = request.args.get("format")
    scope = request.args.get("scope", "dashboard")
    if file_format not in {"xlsx", "pdf"}:
        return jsonify({"message": "invalid_format"}), 400
    if scope not in {"dashboard", "dashboard-and-transactions"}:
        return jsonify({"message": "invalid_scope"}), 400

    try:
        start_date = _parse_date(request.args.get("start_date"))
        end_date = _parse_date(request.args.get("end_date"))
    except ValueError:
        return jsonify({"message": "invalid_date"}), 400
    if start_date and end_date and start_date > end_date:
        return jsonify({"message": "invalid_date_range"}), 400

    receivables, payables = _load_records(user_id, start_date, end_date)
    include_transactions = scope == "dashboard-and-transactions"
    if file_format == "xlsx":
        file_stream = _create_xlsx(receivables, payables, include_transactions)
        mimetype = EXCEL_MIME_TYPE
    else:
        file_stream = _create_pdf(receivables, payables, include_transactions, start_date, end_date)
        mimetype = "application/pdf"

    return send_file(
        file_stream,
        mimetype=mimetype,
        as_attachment=True,
        download_name=f"walletflow-dashboard.{file_format}",
    )
